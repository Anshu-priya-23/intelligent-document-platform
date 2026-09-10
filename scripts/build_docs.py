"""Create the submission diagram and presentation from editable Python content."""
from pathlib import Path
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT = Path("docs")
OUT.mkdir(exist_ok=True)
slides = [
    ("Document Intelligence Platform", ["Financial document extraction, validation and review", "Scope: invoice, balance sheet, profit & loss, cash flow", "Submission status: locally verified; completeness and deployment pending"]),
    ("Problem and scope", ["Accept PDF, JPG and PNG, including scanned documents, at most 3 pages", "Extract meaningful fields, tables and comparative-period values", "Return null for missing data and evidence for supported values", "Make calculations and stored results inspectable through UI and API"]),
    ("Architecture", ["Browser upload -> FastAPI -> file validation", "PyMuPDF native text / Tesseract OCR -> configurable LLM", "Pydantic and evidence checks -> Decimal financial checks", "SQLAlchemy -> SQLite locally / PostgreSQL in deployment", "Dashboard and document-name API read the same persistent record"]),
    ("Input validation and OCR", ["Check extension, content signature, byte size, image decode and PDF render", "Reject corruption, encryption, repaired PDFs and more than 3 pages", "Native text where adequate; render image-dominated pages for OCR", "EXIF correction and orientation detection; bounded OCR runtime", "Supplied inventory: 30 PDFs and 20 JPEGs; most PDFs need OCR"]),
    ("Structured extraction and grounding", ["User selects document type; no classification service", "Dynamic fields, full table rows and separate comparative periods", "JSON schema plus server-side validation; reject truncated responses", "Source excerpt and page number accompany values", "Clear values lacking evidence; OCR evidence is not proof of accuracy"]),
    ("Financial validation", ["Invoice: quantity x price, line sum, tax-inclusive/exclusive total and cash/change", "Balance sheet: assets vs capital/liabilities and complete component sums", "P&L: banking income, expenditure, minority interest and appropriation", "Cash flow: activities plus FX; opening plus net change and adjustments", "Decimal arithmetic; absolute 0.02 tolerance; brackets are negative", "Missing operands -> NOT_APPLICABLE; each period checked separately"]),
    ("API, persistence and frontend", ["POST /api/v1/documents/process; GET /api/v1/documents/{document_name}", "GET /api/v1/documents; GET /api/v1/health; Swagger at /docs", "Atomic upsert: latest completed result per document name", "Searchable database dashboard, fields, tables, validation and raw JSON", "Clear missing/failure highlights; controlled error messages"]),
    ("Testing and actual verification", ["Automated file, arithmetic, grounding, provider-error and API tests", "Full API test uses real PDF/SQLite and explicitly mocked LLM", "Local HTTP: frontend/assets, docs, health and listing verified", "Stored supplied results: all four types, Gemini 2.5 Flash with local OCR", "Sample outputs distinguish stored results, historical errors and synthetic checks", "Browser desktop/mobile verified; no deployed result is claimed"]),
    ("Deployment plan", ["One Render Python service hosts API and frontend; no Docker", "Environment-only provider key, URL, model and PostgreSQL connection", "Build installs Python dependencies and local Tesseract Debian packages", "Linux package resolution and PostgreSQL require deployment validation", "Publishing and deployment deferred; all public URL entries pending"]),
    ("Limitations and next steps", ["Resolve missing statement tables and documented cash-flow OCR error", "Run supplied samples and review every field, table and equation", "OCR errors, unfamiliar layouts and provider output limits remain risks", "Add production authentication, retention, limits, backups and evaluation", "Candidate: review code, run tests, rehearse demo and approve submission"]),
]

# Editable PowerPoint and matching standalone PDF.
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
pdf = pymupdf.open()
for index, (title, bullets) in enumerate(slides, 1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(.6), Inches(.45), Inches(12), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(30)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(23, 60, 75)
    box = slide.shapes.add_textbox(Inches(.7), Inches(1.7), Inches(12), Inches(5))
    frame = box.text_frame
    frame.word_wrap = True
    for i, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.font.size = Pt(22)
        paragraph.space_after = Pt(18)
    footer = slide.shapes.add_textbox(Inches(12), Inches(7), Inches(1), Inches(.3))
    footer.text_frame.text = str(index)
    page = pdf.new_page(width=960, height=540)
    page.draw_rect(pymupdf.Rect(0, 0, 960, 10), color=None, fill=(.09,.24,.29))
    page.insert_text((40,65), title, fontsize=27, color=(.09,.24,.29))
    y = 115
    for bullet in bullets:
        height = 52 if len(bullet) > 92 else 44
        remaining = page.insert_textbox(pymupdf.Rect(48,y,910,y+height), "- " + bullet, fontsize=17, color=(.12,.18,.22))
        if remaining < 0: raise RuntimeError("Slide text does not fit: " + bullet)
        y += height + 12
    page.insert_text((900,515), str(index), fontsize=11)
prs.save(OUT / "solution_presentation.pptx")
pdf.save(OUT / "solution_presentation.pdf")
pdf.close()

# Code-native engineering diagram, also exported as PDF and PNG.
diagram = pymupdf.open()
page = diagram.new_page(width=1000, height=750)
page.insert_text((45,45), "Document Intelligence - Architecture", fontsize=25, color=(.09,.24,.29))
boxes = [
    (pymupdf.Rect(60,90,390,170), "HTML / CSS / JavaScript\nUpload, dashboard, results, raw JSON"),
    (pymupdf.Rect(60,215,390,295), "FastAPI routes and file validation\nPDF/JPG/PNG, size, integrity, <= 3 pages"),
    (pymupdf.Rect(60,340,390,420), "PyMuPDF / Tesseract OCR\nNative text or oriented page images"),
    (pymupdf.Rect(610,340,940,420), "Configurable external LLM API\nText + schema -> grounded structured JSON"),
    (pymupdf.Rect(610,215,940,295), "Pydantic + Decimal validation\nEvidence, fields, periods, financial checks"),
    (pymupdf.Rect(610,90,940,170), "SQLAlchemy persistent result store\nSQLite locally / PostgreSQL deployed"),
]
for rect, label in boxes:
    page.draw_rect(rect, color=(.09,.24,.29), fill=(.93,.96,.97), width=1.3)
    page.insert_textbox(rect + (12,18,-12,-8), label, fontsize=15, align=1)

def arrow(start, end):
    page.draw_line(start, end, color=(.1,.3,.4), width=2)
    x,y=end; dx,dy=end[0]-start[0],end[1]-start[1]
    if dx == 0:
        sign = 1 if dy > 0 else -1
        page.draw_polyline([(x-6,y-sign*10),(x,y),(x+6,y-sign*10)], color=(.1,.3,.4), width=2)
    else:
        sign = 1 if dx > 0 else -1
        page.draw_polyline([(x-sign*10,y-6),(x,y),(x-sign*10,y+6)], color=(.1,.3,.4), width=2)
for a,b in [((225,170),(225,215)),((225,295),(225,340)),((390,380),(610,380)),((775,340),(775,295)),((775,215),(775,170)),((610,130),(390,130))]: arrow(a,b)
page.insert_textbox(pymupdf.Rect(70,490,930,680), "Cross-cutting: environment configuration, controlled exceptions, stage logging and timestamps.\n\nPOST processes synchronously. GET and dashboard read the latest completed database result by filename. Original uploads are not retained.\n\nExternal dependency boundaries: Tesseract executable and provider credentials/model. Local implementation and supplied processing verified; extraction completeness and deployment remain pending.", fontsize=17)
diagram.save(OUT / "architecture.pdf")
page.get_pixmap(matrix=pymupdf.Matrix(1.4,1.4)).save(OUT / "architecture.png")
diagram.close()
print("Created architecture PDF/PNG and presentation PPTX/PDF")
